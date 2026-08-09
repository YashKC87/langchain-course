"""Generate a catalogue-style PowerPoint for the Right Model Lab use case."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "docs" / "Right-Model-Lab-Catalogue.pptx"

# Visual system — slate + emerald (matches lab UI; avoid purple defaults)
NAVY = RGBColor(0x0B, 0x12, 0x20)
SLATE = RGBColor(0x0F, 0x17, 0x2A)
PANEL = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
EMERALD = RGBColor(0x10, 0xB9, 0x81)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
LIGHT = RGBColor(0xE2, 0xE8, 0xF0)


def _set_run(run, size=18, bold=False, color=WHITE, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_bg(slide, color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def _bar(slide, y=0, height=0.12, color=EMERALD):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(y), Inches(13.333), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _textbox(slide, left, top, width, height, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)
    return box


def _bullets(slide, left, top, width, height, lines, size=16, color=LIGHT, spacing=6):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = line
        _set_run(run, size=size, color=color)
    return box


def _card(slide, left, top, width, height, title, body_lines, accent=EMERALD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = RGBColor(0x33, 0x41, 0x55)
    # accent strip
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.08), Inches(height)
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent
    strip.line.fill.background()
    _textbox(slide, left + 0.2, top + 0.12, width - 0.3, 0.35, title, size=15, bold=True, color=WHITE)
    _bullets(slide, left + 0.2, top + 0.48, width - 0.3, height - 0.6, body_lines, size=12, color=LIGHT, spacing=4)


def _section_title(slide, title, subtitle=""):
    _add_bg(slide)
    _bar(slide, 0, 0.1, EMERALD)
    _textbox(slide, 0.6, 0.35, 12, 0.5, title, size=32, bold=True, color=WHITE)
    if subtitle:
        _textbox(slide, 0.6, 0.9, 12, 0.4, subtitle, size=16, color=MUTED)
    _textbox(slide, 0.6, 7.1, 12, 0.25, "Right Model Lab  ·  SLM + Frontier Device Monitoring", size=11, color=MUTED)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Cover
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, 0, 0.15, EMERALD)
    _textbox(s, 0.8, 1.8, 11.5, 0.4, "USE CASE CATALOGUE", size=14, bold=True, color=EMERALD)
    _textbox(s, 0.8, 2.3, 11.5, 0.8, "Right Model Lab", size=44, bold=True, color=WHITE)
    _textbox(
        s,
        0.8,
        3.2,
        11.5,
        0.6,
        "SLM + Frontier Architecture Patterns for Digital Workplace Device Monitoring",
        size=18,
        color=LIGHT,
    )
    _textbox(
        s,
        0.8,
        4.1,
        11.5,
        0.5,
        "Principle: Use the right model for the right task.",
        size=20,
        bold=True,
        color=CYAN,
    )
    _bullets(
        s,
        0.8,
        5.0,
        11,
        1.5,
        [
            "Five locked architecture patterns  ·  One scenario each",
            "Measurable tokens · latency · confidence · illustrative cost",
            "UI: Overview  ·  AI Model Operations  ·  Model Comparison",
        ],
        size=15,
        color=MUTED,
    )

    # 2 Catalogue contents
    s = prs.slides.add_slide(blank)
    _section_title(s, "Catalogue Contents", "Single deck for the complete use case")
    items = [
        ("01", "Business problem & principle"),
        ("02", "Solution architecture"),
        ("03", "SLM vs Frontier roles"),
        ("04", "Five patterns at a glance"),
        ("05–09", "Pattern catalogue (detail)"),
        ("10", "Data, models & observability"),
        ("11", "UI catalogue — three dropdown pages"),
        ("12", "Primary metrics & provenance"),
        ("13", "Demo path & proof points"),
        ("14", "How to run"),
    ]
    for i, (num, text) in enumerate(items):
        col = i // 5
        row = i % 5
        left = 0.8 + col * 6.2
        top = 1.6 + row * 0.85
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(5.8), Inches(0.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PANEL
        shape.line.color.rgb = RGBColor(0x33, 0x41, 0x55)
        _textbox(s, left + 0.25, top + 0.18, 1.2, 0.4, num, size=16, bold=True, color=EMERALD)
        _textbox(s, left + 1.5, top + 0.18, 4.0, 0.4, text, size=16, color=WHITE)

    # 3 Problem
    s = prs.slides.add_slide(blank)
    _section_title(s, "01  ·  Business Problem", "Digital workplace monitoring workloads are mixed")
    _card(
        s,
        0.6,
        1.5,
        5.9,
        4.8,
        "Routine high-volume work",
        [
            "Health / status checks",
            "Known SOP remediation",
            "Crash counts & simple classification",
            "Risk if sent to Frontier: wasted capacity & cost",
        ],
        accent=EMERALD,
    )
    _card(
        s,
        6.8,
        1.5,
        5.9,
        4.8,
        "Hard reasoning work",
        [
            "Multi-domain RCA",
            "Ambiguous Teams + VPN symptoms",
            "Planning and final synthesis",
            "Risk if sent only to SLM: quality collapse",
        ],
        accent=BLUE,
    )

    # 4 Principle / solution
    s = prs.slides.add_slide(blank)
    _section_title(s, "01  ·  Principle & Solution", "Hybrid collaboration — not SLM replacement")
    _textbox(s, 0.8, 1.6, 11.5, 0.6, "Use the right model for the right task.", size=28, bold=True, color=CYAN)
    _bullets(
        s,
        0.8,
        2.5,
        11.5,
        4,
        [
            "Reserve Frontier for planning, ambiguity resolution, and synthesis.",
            "Use SLMs for routine, repetitive, high-volume steps.",
            "Measure success as Frontier tokens avoided — with quality still visible.",
            "Never present ESTIMATED baselines as actual spend.",
            "One locked scenario per pattern keeps the demo explainable.",
        ],
        size=18,
        color=LIGHT,
        spacing=10,
    )

    # 5 Architecture
    s = prs.slides.add_slide(blank)
    _section_title(s, "02  ·  Solution Architecture", "Layers in the lab")
    layers = [
        ("UI", "Streamlit — Right Model Lab\nOverview · AI Model Operations · Model Comparison", CYAN),
        ("Orchestration", "ScenarioService · FastAPI\nRun one / run all · ResultStore", EMERALD),
        ("Patterns", "Planner-Worker · Router · Cascade · RAG · Fallback", BLUE),
        ("Models + Observability", "LanguageModel interface · Demo / Ollama / Azure\nTraceStore · LangSmith · BaselineEstimator", AMBER),
    ]
    for i, (title, body, accent) in enumerate(layers):
        top = 1.5 + i * 1.25
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(top), Inches(10.8), Inches(1.1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PANEL
        shape.line.color.rgb = RGBColor(0x33, 0x41, 0x55)
        strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(top), Inches(0.12), Inches(1.1))
        strip.fill.solid()
        strip.fill.fore_color.rgb = accent
        strip.line.fill.background()
        _textbox(s, 1.6, top + 0.15, 3, 0.35, title, size=18, bold=True, color=WHITE)
        _textbox(s, 4.8, top + 0.15, 6.8, 0.8, body, size=14, color=LIGHT)

    # 6 Roles
    s = prs.slides.add_slide(blank)
    _section_title(s, "03  ·  SLM vs Frontier", "Executive model analogy")
    _card(
        s,
        0.6,
        1.5,
        5.9,
        5.0,
        "SLM — Service Desk Engineer",
        [
            "Fast and lower cost",
            "Strong on known / repetitive work",
            "Classification, extraction, telemetry interpretation",
            "Routine summaries and grounded SOP answers",
            "Providers in lab: Demo or Ollama-compatible local model",
        ],
        accent=EMERALD,
    )
    _card(
        s,
        6.8,
        1.5,
        5.9,
        5.0,
        "Frontier — Senior Solution Architect",
        [
            "Stronger reasoning and planning",
            "Better at ambiguity and cross-domain synthesis",
            "More expensive / often slower",
            "Best reserved for high-value segments",
            "Providers in lab: Demo or Azure OpenAI-compatible",
        ],
        accent=BLUE,
    )

    # 7 Patterns glance
    s = prs.slides.add_slide(blank)
    _section_title(s, "04  ·  Five Patterns at a Glance", "One locked scenario each")
    rows = [
        ("Planner-Worker", "LAPTOP-1204 RCA", "Selective reasoning", EMERALD),
        ("Router", "100 health requests", "High-volume efficiency", CYAN),
        ("Confidence Cascade", "LAPTOP-9910 Teams/VPN", "Quality-aware escalation", BLUE),
        ("RAG", "LAPTOP-1507 disk/OneDrive", "Grounded smaller model", AMBER),
        ("Fallback", "Frontier unavailable", "Resilience / availability", RGBColor(0x38, 0xBD, 0xF8)),
    ]
    # header
    for i, h in enumerate(["Pattern", "Locked scenario", "Primary benefit"]):
        _textbox(s, 0.7 + i * 4.0, 1.5, 3.8, 0.35, h, size=13, bold=True, color=MUTED)
    for r, (p, sc, b, accent) in enumerate(rows):
        top = 2.0 + r * 0.9
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(top), Inches(12.1), Inches(0.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PANEL
        shape.line.color.rgb = RGBColor(0x33, 0x41, 0x55)
        strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(top), Inches(0.1), Inches(0.8))
        strip.fill.solid()
        strip.fill.fore_color.rgb = accent
        strip.line.fill.background()
        _textbox(s, 0.9, top + 0.22, 3.6, 0.4, p, size=16, bold=True, color=WHITE)
        _textbox(s, 4.7, top + 0.22, 3.8, 0.4, sc, size=15, color=LIGHT)
        _textbox(s, 8.7, top + 0.22, 3.8, 0.4, b, size=15, color=EMERALD)

    # Pattern detail slides
    patterns = [
        (
            "05  ·  Planner-Worker Hierarchy",
            "Selective reasoning + orchestration efficiency",
            [
                "Device: LAPTOP-1204 — multi-domain RCA",
                "Flow: Frontier Plan → SLM Workers → Frontier Synthesis",
                "SLM workers: CPU · Teams · Network/VPN · Compliance · History",
                "Frontier only where orchestration and judgment matter",
                "Proof: Frontier tokens avoided vs ESTIMATED all-Frontier",
            ],
            "Incident → FRONTIER Planner → SLM Workers → FRONTIER RCA",
        ),
        (
            "06  ·  Router Pattern",
            "High-volume inference efficiency",
            [
                "Volume: 100 seeded endpoint health requests",
                "NON-LLM router classifies routine vs complex",
                "Demo mix: ~85% SLM · ~15% Frontier",
                "Routine: status / counts / simple classification",
                "Complex: cross-domain RCA · ambiguous telemetry",
            ],
            "Request → NON-LLM Router → SLM (routine) | FRONTIER (complex)",
        ),
        (
            "07  ·  Confidence Cascade",
            "Quality-aware escalation",
            [
                "Device: LAPTOP-9910 — ambiguous Teams + VPN",
                "SLM first-pass diagnosis → confidence gate",
                "Threshold: CONFIDENCE_THRESHOLD = 0.85",
                "Demo path: SLM ~0.54 → escalate to Frontier",
                "Honest note: one escalated case can cost more than direct Frontier; fleet economics still win",
            ],
            "Incident → SLM → Gate → Accept | FRONTIER escalate",
        ),
        (
            "08  ·  RAG",
            "Grounded enterprise knowledge with a smaller model",
            [
                "Device: LAPTOP-1507 — disk + OneDrive SOP remediation",
                "NON-LLM retrieve / rank / context from Markdown SOP KB",
                "SLM grounded generation with citations",
                "Frontier not required for this scenario",
                "Proof: grounding + RAG+SLM vs ESTIMATED RAG+Frontier",
            ],
            "Ticket → Retrieve SOP → Rank → SLM grounded answer",
        ),
        (
            "09  ·  Designing for Fallback",
            "Resilience and availability — not token optimization",
            [
                "Device: LAPTOP-1204 during Frontier unavailability",
                "Failure modes: 503 · timeout · rate limit · auth · offline · cost guard",
                "Flow: Frontier attempt → failure detector → SLM fallback",
                "Degraded quality (~0.72 confidence in demo) with continuity",
                "Proof: operations continue when Frontier is down",
            ],
            "Incident → FRONTIER → Fail → SLM fallback → Response",
        ),
    ]
    for title, subtitle, bullets, edge in patterns:
        s = prs.slides.add_slide(blank)
        _section_title(s, title, subtitle)
        _card(s, 0.6, 1.5, 7.8, 5.0, "What this pattern demonstrates", bullets, accent=EMERALD)
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.7), Inches(1.5), Inches(4.0), Inches(5.0))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PANEL
        shape.line.color.rgb = RGBColor(0x33, 0x41, 0x55)
        _textbox(s, 8.95, 1.7, 3.6, 0.4, "Architecture edge", size=14, bold=True, color=CYAN)
        _textbox(s, 8.95, 2.3, 3.6, 3.5, edge, size=15, color=LIGHT)

    # 10 Data / models / obs
    s = prs.slides.add_slide(blank)
    _section_title(s, "10  ·  Data, Models & Observability", "What the lab is built on")
    _card(
        s,
        0.5,
        1.5,
        4.0,
        5.0,
        "Synthetic data",
        [
            "100-device fleet (seeded)",
            "Incidents + DEX / risk scores",
            "Locked devices: 1204 · 9910 · 1507",
            "Markdown SOP knowledge base",
            "TF-IDF VectorStore for RAG",
        ],
        accent=EMERALD,
    )
    _card(
        s,
        4.7,
        1.5,
        4.0,
        5.0,
        "Model layer",
        [
            "Shared LanguageModel interface",
            "Demo clients (default)",
            "Ollama-compatible local SLM",
            "Azure OpenAI-compatible Frontier",
            "Patterns never import SDKs directly",
        ],
        accent=BLUE,
    )
    _card(
        s,
        8.9,
        1.5,
        4.0,
        5.0,
        "Observability",
        [
            "TokenMetrics on every segment",
            "Local TraceStore always on",
            "Optional LangSmith remote",
            "BaselineEstimator (ESTIMATED)",
            "Illustrative cost calculator",
        ],
        accent=AMBER,
    )

    # 11 UI catalogue
    s = prs.slides.add_slide(blank)
    _section_title(s, "11  ·  UI Catalogue", "Three dropdown pages")
    _card(
        s,
        0.5,
        1.5,
        4.0,
        5.0,
        "Overview",
        [
            "Mode · models · fleet strip",
            "KPI row & Run all",
            "Five pattern tiles",
            "Open → architecture diagram",
            "Scenario → Flow → Why → Tokens",
            "Trace · Tail · Architect Decision",
        ],
        accent=EMERALD,
    )
    _card(
        s,
        4.7,
        1.5,
        4.0,
        5.0,
        "AI Model Operations",
        [
            "MLOps control tower",
            "17 KPIs incl. avoidance",
            "12 interactive charts + radar",
            "Selected vs all-Frontier",
            "Avoidance rate · cost vs quality",
            "Pattern efficiency table",
        ],
        accent=CYAN,
    )
    _card(
        s,
        8.9,
        1.5,
        4.0,
        5.0,
        "Model Comparison",
        [
            "Segment-level decision audit",
            "Selected vs ESTIMATED alternative",
            "Tokens · latency · confidence",
            "Decision + Why columns",
            "Per-pattern avoidance callout",
            "Defends each model choice",
        ],
        accent=BLUE,
    )

    # 12 Metrics
    s = prs.slides.add_slide(blank)
    _section_title(s, "12  ·  Primary Metrics & Provenance", "How claims stay honest")
    _card(
        s,
        0.6,
        1.5,
        6.0,
        5.0,
        "Primary metric",
        [
            "Frontier tokens avoided",
            "Hybrid vs ESTIMATED all-Frontier baseline",
            "Also track: Frontier calls avoided",
            "Supporting: latency · confidence · illustrative cost",
            "Fewer total tokens ≠ automatically better",
        ],
        accent=EMERALD,
    )
    _card(
        s,
        6.9,
        1.5,
        5.8,
        5.0,
        "Provenance labels",
        [
            "ACTUAL — provider-reported usage",
            "SIMULATED — Demo Mode telemetry",
            "ESTIMATED — baselines / alternatives only",
            "Never present ESTIMATED as money spent",
            "Baseline does not call Frontier just for the chart",
        ],
        accent=AMBER,
    )

    # 13 Demo path
    s = prs.slides.add_slide(blank)
    _section_title(s, "13  ·  Demo Path & Proof Points", "Recommended walkthrough")
    steps = [
        ("1", "Overview — status, KPIs, name five tiles"),
        ("2", "Open Planner-Worker — architecture, Load result, Architect Decision"),
        ("3", "Optional: Cascade (quality) or Fallback (resilience)"),
        ("4", "AI Model Operations — charts 2, 9, 11 + efficiency table"),
        ("5", "Model Comparison — one SLM win row + one Frontier win row"),
        ("6", "Close on principle + Frontier tokens avoided"),
    ]
    for i, (n, text) in enumerate(steps):
        top = 1.5 + i * 0.85
        circle = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(top), Inches(0.55), Inches(0.55))
        circle.fill.solid()
        circle.fill.fore_color.rgb = EMERALD
        circle.line.fill.background()
        _textbox(s, 0.8, top + 0.1, 0.55, 0.4, n, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), Inches(top), Inches(10.8), Inches(0.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PANEL
        shape.line.color.rgb = RGBColor(0x33, 0x41, 0x55)
        _textbox(s, 1.9, top + 0.18, 10.2, 0.4, text, size=16, color=WHITE)

    # 14 How to run
    s = prs.slides.add_slide(blank)
    _section_title(s, "14  ·  How to Run", "Demo Mode by default — Live optional")
    _bullets(
        s,
        0.8,
        1.6,
        11.5,
        5,
        [
            "python3 -m venv .venv && source .venv/bin/activate",
            "pip install -r requirements.txt && cp .env.example .env",
            "streamlit run app/ui/dashboard.py   →  http://localhost:8501",
            "Optional API: uvicorn app.main:app --reload --port 8000",
            "DEMO_MODE=true for offline deterministic demos",
            "Live: Ollama SLM + Azure OpenAI Frontier + optional LangSmith",
            "Docs: README.md · docs/copy-paste-script.md · docs/use-case-demonstration.md",
        ],
        size=17,
        color=LIGHT,
        spacing=10,
    )

    # 15 Close
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, 0, 0.15, EMERALD)
    _textbox(s, 0.8, 2.2, 11.5, 0.5, "Catalogue summary", size=16, bold=True, color=EMERALD)
    _textbox(s, 0.8, 2.8, 11.5, 0.8, "Right Model Lab", size=40, bold=True, color=WHITE)
    _textbox(
        s,
        0.8,
        3.7,
        11.5,
        0.6,
        "Right model  ·  Right task  ·  Measurable evidence",
        size=22,
        color=CYAN,
    )
    _bullets(
        s,
        0.8,
        4.6,
        11.5,
        2,
        [
            "Five patterns · three UI pages · honest provenance · Frontier tokens avoided",
            "Built for digital workplace device monitoring architecture demos",
        ],
        size=16,
        color=MUTED,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
